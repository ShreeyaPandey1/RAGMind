import os
from typing import List, Dict, Any, Optional
import fitz # PyMuPDF
import docx
from app.core.exceptions import DocumentProcessingError, UnsupportedFileTypeError
from app.core.logging import logger

class ParsedPage:
    def __init__(self, page_number: int, text: str, metadata: Optional[Dict[str, Any]] = None):
        self.page_number = page_number
        self.text = text.strip()
        self.metadata = metadata or {}

class ParsedDocument:
    def __init__(self, pages: List[ParsedPage], total_pages: int, mime_type: str):
        self.pages = pages
        self.total_pages = total_pages
        self.mime_type = mime_type

class DocumentParserService:
    @staticmethod
    def parse_file(file_path: str, mime_type: str, original_filename: str) -> ParsedDocument:
        ext = os.path.splitext(original_filename)[1].lower()
        logger.info(f"Parsing document '{original_filename}' (ext: {ext}, mime: {mime_type})")

        if ext == ".pdf":
            return DocumentParserService._parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return DocumentParserService._parse_docx(file_path)
        elif ext in [".txt", ".md", ".markdown"]:
            return DocumentParserService._parse_text(file_path, ext)
        else:
            raise UnsupportedFileTypeError(f"Unsupported document format: {ext}")

    @staticmethod
    def _parse_pdf(file_path: str) -> ParsedDocument:
        try:
            doc = fitz.open(file_path)
        except Exception as e:
            raise DocumentProcessingError(f"Corrupt or invalid PDF file: {str(e)}")

        if doc.is_encrypted:
            try:
                # Try decrypting with empty password
                doc.authenticate("")
            except Exception:
                raise DocumentProcessingError("PDF is password protected / encrypted and cannot be read.")

        total_pages = len(doc)
        pages: List[ParsedPage] = []
        total_extracted_text = 0

        for page_idx in range(total_pages):
            page = doc[page_idx]
            blocks = page.get_text("blocks")
            cleaned_blocks = []
            for b in blocks:
                # b = (x0, y0, x1, y1, text, block_no, block_type)
                if len(b) >= 7 and b[6] == 0 and b[4]:
                    raw_block = b[4].strip()
                    if raw_block:
                        lines = [line.strip() for line in raw_block.splitlines() if line.strip()]
                        block_text = " ".join(lines)
                        cleaned_blocks.append(block_text)

            clean_text = "\n\n".join(cleaned_blocks)
            if not clean_text:
                fallback = page.get_text("text") or ""
                clean_text = "\n".join(l.strip() for l in fallback.splitlines() if l.strip())

            total_extracted_text += len(clean_text)
            pages.append(ParsedPage(
                page_number=page_idx + 1,
                text=clean_text,
                metadata={"page_width": page.rect.width, "page_height": page.rect.height}
            ))

        doc.close()

        if total_extracted_text == 0:
            raise DocumentProcessingError(
                "No readable text found in PDF. The document may be scanned or image-based, which requires OCR."
            )

        return ParsedDocument(pages=pages, total_pages=total_pages, mime_type="application/pdf")

    @staticmethod
    def _parse_docx(file_path: str) -> ParsedDocument:
        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise DocumentProcessingError(f"Corrupt or invalid DOCX document: {str(e)}")

        paragraphs = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        if not full_text.strip():
            raise DocumentProcessingError("DOCX file contains no extractable text.")

        return ParsedDocument(
            pages=[ParsedPage(page_number=1, text=full_text, metadata={"paragraph_count": len(paragraphs)})],
            total_pages=1,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    @staticmethod
    def _parse_text(file_path: str, ext: str) -> ParsedDocument:
        encodings = ["utf-8", "latin-1", "cp1252"]
        content = None
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                    break
            except UnicodeDecodeError:
                continue

        if content is None:
            raise DocumentProcessingError("Unable to decode text file. Unsupported encoding.")

        if not content.strip():
            raise DocumentProcessingError("Text document is empty.")

        mime_type = "text/markdown" if ext in [".md", ".markdown"] else "text/plain"
        return ParsedDocument(
            pages=[ParsedPage(page_number=1, text=content, metadata={"char_count": len(content)})],
            total_pages=1,
            mime_type=mime_type
        )
